from __future__ import annotations

import hashlib
import json
import math
import secrets
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import delete, select

from .audit import add_audit, audit_event_digest, backfill_audit_chain
from .config import INSTANCE_DIR, PROJECT_DIR, settings
from .db.base import Base, ENGINE, SessionLocal, UPLOAD_DIR
from .db.models import *  # noqa: F401,F403 - bootstrap consumes the canonical ORM models
from .database import (
    _ensure_v020_defaults,
    _ensure_v021_defaults,
    _ensure_v022_defaults,
    _ensure_v023_defaults,
    _ensure_v024_defaults,
    _ensure_v025_defaults,
    _ensure_v026_defaults,
    _ensure_v027_defaults,
    _ensure_v028_defaults,
    _ensure_v030_defaults,
    _ensure_v031_defaults,
    _ensure_v032_defaults,
    _ensure_v033_defaults,
    _ensure_v034_defaults,
    _ensure_v035_defaults,
    _ensure_v036_defaults,
    _ensure_v037_defaults,
    _ensure_v043_defaults,
    _ensure_v044_defaults,
    _ensure_v045_defaults,
    _ensure_v050_defaults,
    _ensure_v100_final_defaults,
    _seed_methodology,
    _seed_sector_templates,
    hash_password,
    refresh_progress,
    write_simple_pdf,
)

def init_db() -> None:
    Base.metadata.create_all(ENGINE)
    with SessionLocal() as session:
        existing = session.scalar(select(Organization).limit(1))
        if existing:
            _ensure_v020_defaults(session)
            _ensure_v021_defaults(session)
            _ensure_v022_defaults(session)
            _ensure_v023_defaults(session)
            _ensure_v024_defaults(session)
            _ensure_v025_defaults(session)
            _ensure_v026_defaults(session)
            _ensure_v027_defaults(session)
            _ensure_v028_defaults(session)
            _ensure_v030_defaults(session)
            _ensure_v031_defaults(session)
            _ensure_v032_defaults(session)
            _ensure_v033_defaults(session)
            _ensure_v034_defaults(session)
            _ensure_v035_defaults(session)
            _ensure_v036_defaults(session)
            _ensure_v037_defaults(session)
            _ensure_v043_defaults(session)
            _ensure_v044_defaults(session)
            _ensure_v045_defaults(session)
            _ensure_v050_defaults(session)
            _ensure_v100_final_defaults(session)
            session.commit()
            return

        methodology = _seed_methodology(session)
        _seed_sector_templates(session)
        versions: dict[str, EmissionFactorVersion] = methodology["versions"]  # type: ignore[assignment]

        if not settings.seed_demo:
            if settings.bootstrap_admin_email and settings.bootstrap_admin_password:
                org = Organization(
                    name=settings.bootstrap_organization, trade_name=settings.bootstrap_organization,
                    tax_id="PENDIENTE", sector="Por configurar", country="Colombia", city="Por configurar",
                    contact_name="Administrador", contact_email=settings.bootstrap_admin_email, status="Activa",
                )
                session.add(org)
                session.flush()
                session.add(AppUser(
                    organization_id=org.id, email=settings.bootstrap_admin_email, name="Administrador inicial",
                    role="Administrador", password_hash=hash_password(settings.bootstrap_admin_password), active=True,
                ))
            _ensure_v020_defaults(session)
            _ensure_v021_defaults(session)
            _ensure_v022_defaults(session)
            _ensure_v023_defaults(session)
            _ensure_v024_defaults(session)
            _ensure_v025_defaults(session)
            _ensure_v026_defaults(session)
            _ensure_v027_defaults(session)
            _ensure_v028_defaults(session)
            _ensure_v030_defaults(session)
            _ensure_v031_defaults(session)
            _ensure_v032_defaults(session)
            _ensure_v033_defaults(session)
            _ensure_v034_defaults(session)
            _ensure_v035_defaults(session)
            _ensure_v036_defaults(session)
            _ensure_v037_defaults(session)
            _ensure_v043_defaults(session)
            _ensure_v044_defaults(session)
            _ensure_v045_defaults(session)
            _ensure_v050_defaults(session)
            _ensure_v100_final_defaults(session)
            session.commit()
            return

        org = Organization(
            name="Industrias Andinas Demo S.A.S.", trade_name="Industrias Andinas", tax_id="901.555.101-8",
            sector="Manufactura", ciiu_code="C2029", country="Colombia", department="Antioquia", city="Medellín",
            employees=186, annual_revenue=18_500_000_000, contact_name="Ana Martínez", contact_email="ambiental@industriasandinas.demo",
        )
        session.add(org)
        session.flush()

        demo_password = hash_password("Demo2026!")
        session.add_all([
            AppUser(organization_id=org.id, email="admin@calculatuhuella.local", name="Laura Méndez", role="Administrador", password_hash=demo_password),
            AppUser(organization_id=org.id, email="consultor@calculatuhuella.local", name="Carlos Uribe", role="Consultor", password_hash=demo_password),
            AppUser(organization_id=org.id, email="cliente@calculatuhuella.local", name="Ana Martínez", role="Cliente", password_hash=demo_password),
            AppUser(organization_id=org.id, email="revisor@calculatuhuella.local", name="María Fernández", role="Revisor", password_hash=demo_password),
            AppUser(organization_id=org.id, email="verificador@calculatuhuella.local", name="Andrés Salazar", role="Verificador", password_hash=demo_password),
        ])

        session.flush()
        demo_users = {user.role: user for user in session.scalars(select(AppUser).where(AppUser.organization_id == org.id))}
        for demo_user in demo_users.values():
            session.add(NotificationPreference(user_id=demo_user.id, in_app_enabled=True, email_enabled=True, digest_frequency="Inmediato"))
        session.add_all([
            Notification(organization_id=org.id, user_id=demo_users["Consultor"].id, title="Divulgación climática disponible", message="La V0.20 incorporó comparación de escenarios, matriz de divulgación y paquete ejecutivo para comité directivo.", link="/divulgacion-climatica", category="Producto", priority="Normal", status="Entregada"),
            Notification(organization_id=org.id, user_id=demo_users["Revisor"].id, title="Revisión pendiente", message="El inventario corporativo 2025 contiene observaciones abiertas que requieren seguimiento.", link="/control", category="Revisión", priority="Alta", status="Entregada"),
            Notification(organization_id=org.id, user_id=demo_users["Administrador"].id, title="Configuración productiva", message="Revisa almacenamiento, correo y migraciones desde Administración de plataforma.", link="/administracion-plataforma", category="Operación", priority="Normal", status="Entregada"),
        ])
        session.add_all([
            PlatformSetting(organization_id=org.id, key="brand_descriptor", value="Plataforma profesional de huella de carbono", description="Descriptor visible de la marca"),
            PlatformSetting(organization_id=org.id, key="default_methodology", value="GHG Protocol + ISO 14064-1", description="Metodología sugerida para inventarios nuevos"),
            PlatformSetting(organization_id=org.id, key="data_retention_years", value="7", value_type="integer", description="Retención documental recomendada"),
        ])

        facilities = [
            Facility(organization_id=org.id, name="Planta Medellín", facility_type="Planta de producción", city="Medellín", address="Zona industrial Guayabal", employees=118),
            Facility(organization_id=org.id, name="Bodega Rionegro", facility_type="Centro de distribución", city="Rionegro", address="Zona Franca", employees=42),
            Facility(organization_id=org.id, name="Oficina administrativa", facility_type="Oficina", city="Medellín", address="El Poblado", employees=26),
        ]
        session.add_all(facilities)
        session.flush()

        inv = Inventory(
            organization_id=org.id, name="Inventario corporativo 2025", start_date=date(2025, 1, 1), end_date=date(2025, 12, 31),
            objective="Establecer la línea base corporativa, responder requisitos de clientes y priorizar reducciones.", base_year=2025,
            methodology="GHG Protocol + ISO 14064-1", methodology_version="GHG Protocol Corporate Standard · ISO 14064-1:2018",
            gwp_version="IPCC AR6 · 100 años", consolidation_approach="Control operacional", materiality_threshold=5,
            status="Activo", progress=0, current_stage="Cálculo", notes="Inventario demostrativo con motor matemático, control profesional, cadena de valor, escenarios y verificación y operación productiva, migraciones, almacenamiento, notificaciones, automatizaciones e integraciones, gobierno metodológico, cumplimiento y operación SaaS, flujo comercial y operación contractual y éxito del cliente, inteligencia de impacto y riesgos climáticos V0.20.", version="0.20",
        )
        session.add(inv)
        session.flush()
        for facility in facilities:
            session.add(InventoryFacility(inventory_id=inv.id, facility_id=facility.id, included=True, inclusion_percentage=100))

        sources = [
            EmissionSource(inventory_id=inv.id, facility_id=facilities[0].id, name="Electricidad", scope=2, category="Energía adquirida", responsible="Contabilidad", materiality="Alta", data_frequency="Mensual", preferred_unit="kWh", icon="bolt"),
            EmissionSource(inventory_id=inv.id, facility_id=facilities[0].id, name="Diésel", scope=1, category="Combustión fija", responsible="Mantenimiento", materiality="Alta", data_frequency="Mensual", preferred_unit="L", icon="fuel"),
            EmissionSource(inventory_id=inv.id, facility_id=facilities[1].id, name="Vehículos", scope=1, category="Combustión móvil", responsible="Logística", materiality="Alta", data_frequency="Mensual", preferred_unit="L", icon="truck"),
            EmissionSource(inventory_id=inv.id, facility_id=facilities[0].id, name="Refrigerantes", scope=1, category="Emisiones fugitivas", responsible="Mantenimiento", materiality="Media", data_frequency="Anual", preferred_unit="kg", icon="snow"),
            EmissionSource(inventory_id=inv.id, facility_id=facilities[0].id, name="Residuos", scope=3, category="Residuos operacionales", responsible="Gestión ambiental", materiality="Media", data_frequency="Mensual", preferred_unit="t", icon="waste"),
            EmissionSource(inventory_id=inv.id, facility_id=facilities[1].id, name="Transporte contratado", scope=3, category="Transporte y distribución", responsible="Logística", materiality="Alta", data_frequency="Mensual", preferred_unit="t·km", icon="route"),
            EmissionSource(inventory_id=inv.id, facility_id=None, name="Bienes y servicios adquiridos", scope=3, category="Datos específicos de proveedores", responsible="Compras sostenibles", materiality="Alta", data_frequency="Anual", preferred_unit="tCO₂e", icon="suppliers"),
        ]
        session.add_all(sources)
        session.flush()

        assignment_map = {
            0: ["Electricidad de red Colombia · demo"],
            1: ["Diésel combustión fija · CO2 demo", "Diésel combustión fija · CH4 demo", "Diésel combustión fija · N2O demo"],
            2: ["Gasolina vehículos · CO2e demo"],
            3: ["Refrigerante · demo"],
            4: ["Residuos gestionados · demo"],
            5: ["Transporte de carga · demo"],
        }
        for index, names in assignment_map.items():
            for factor_name in names:
                session.add(SourceFactorAssignment(source_id=sources[index].id, factor_version_id=versions[factor_name].id, active=True, assigned_by="sistema", notes="Asignación demostrativa V0.4"))

        demo_dir = UPLOAD_DIR / f"org_{org.id}" / f"inventory_{inv.id}"
        demo_dir.mkdir(parents=True, exist_ok=True)
        pdf_path = demo_dir / "facturas_energia_2025.pdf"
        write_simple_pdf(pdf_path, "Calcula tu Huella - soporte demostrativo de electricidad 2025")
        pdf_bytes = pdf_path.read_bytes()
        electricity_doc = EvidenceDocument(
            inventory_id=inv.id, source_id=sources[0].id, name="Facturas_energia_2025.pdf",
            stored_name=str(pdf_path.relative_to(INSTANCE_DIR)), document_type="Factura", source_name="Electricidad",
            period_label="Enero–diciembre 2025", status="Aprobado", uploaded_by="Ana Martínez",
            file_size=len(pdf_bytes), sha256=hashlib.sha256(pdf_bytes).hexdigest(), notes="Soporte demostrativo descargable.",
        )
        session.add(electricity_doc)

        xlsx_path = demo_dir / "consolidado_diesel_2025.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Consolidado"
        ws.append(["Mes", "Litros"])
        diesel_values = [1210, 1185, 1250, 1190, 1280, 1310, 1275, 1235, 1295, 1325, 1260, 1305]
        for month, value in enumerate(diesel_values, 1):
            ws.append([date(2025, month, 1), value])
        wb.save(xlsx_path)
        xlsx_bytes = xlsx_path.read_bytes()
        diesel_doc = EvidenceDocument(
            inventory_id=inv.id, source_id=sources[1].id, name="Consolidado_diesel_2025.xlsx",
            stored_name=str(xlsx_path.relative_to(INSTANCE_DIR)), document_type="Registro operativo", source_name="Diésel",
            period_label="2025", status="En revisión", uploaded_by="Ana Martínez", file_size=len(xlsx_bytes),
            sha256=hashlib.sha256(xlsx_bytes).hexdigest(),
        )
        session.add(diesel_doc)
        session.flush()

        electricity_values = [18450, 17980, 18820, 18140, 19010, 19480, 19220, 18790, 19610, 19930, 19350, 20120]
        for month, value in enumerate(electricity_values, 1):
            session.add(ActivityData(source_id=sources[0].id, evidence_id=electricity_doc.id, period_start=date(2025, month, 1), period_end=date(2025, month, 28), value=value, unit="kWh", data_origin="Factura", quality_level="A", status="Aprobado", notes="Consumo facturado mensual.", created_by="cliente@calculatuhuella.local"))
        for month, value in enumerate(diesel_values, 1):
            session.add(ActivityData(source_id=sources[1].id, evidence_id=diesel_doc.id, period_start=date(2025, month, 1), period_end=date(2025, month, 28), value=value, unit="L", data_origin="Registro operativo", quality_level="B", status="En revisión", notes="Consolidado de abastecimiento.", created_by="cliente@calculatuhuella.local"))
        for month, value in enumerate([780, 760, 810, 795, 820, 835, 800, 790, 845, 830], 1):
            session.add(ActivityData(source_id=sources[2].id, period_start=date(2025, month, 1), period_end=date(2025, month, 28), value=value, unit="L", data_origin="Registro operativo", quality_level="B", status="Cargado", created_by="cliente@calculatuhuella.local"))
        session.add(ActivityData(source_id=sources[3].id, period_start=date(2025, 1, 1), period_end=date(2025, 12, 31), value=34.5, unit="kg", data_origin="Registro operativo", quality_level="B", status="En revisión", created_by="cliente@calculatuhuella.local"))
        for month, value in enumerate([18.2, 17.6, 19.1, 18.8, 20.4, 19.7], 1):
            session.add(ActivityData(source_id=sources[4].id, period_start=date(2025, month, 1), period_end=date(2025, month, 28), value=value, unit="t", data_origin="Certificado", quality_level="B", status="Cargado", created_by="cliente@calculatuhuella.local"))
        for month, value in enumerate([12800, 13450, 13100], 1):
            session.add(ActivityData(source_id=sources[5].id, period_start=date(2025, month, 1), period_end=date(2025, month, 28), value=value, unit="t·km", data_origin="Estimación", quality_level="C", is_estimated=True, status="Provisional", notes="Distancia estimada a partir de ruta estándar.", created_by="cliente@calculatuhuella.local"))

        # Indicadores operativos reales para intensidad y comparación.
        production_values = [980, 1010, 1040, 1025, 1080, 1115, 1090, 1065, 1120, 1140, 1105, 1190]
        for month, value in enumerate(production_values, 1):
            session.add(ActivityIndicator(
                inventory_id=inv.id, facility_id=facilities[0].id,
                period_start=date(2025, month, 1), period_end=date(2025, month, 28),
                indicator_type="Producción", value=value, unit="t",
                source_name="Registro de producción", status="Aprobado",
                created_by="cliente@calculatuhuella.local",
            ))
        session.add(ActivityIndicator(
            inventory_id=inv.id, period_start=date(2025, 1, 1), period_end=date(2025, 12, 31),
            indicator_type="Empleados", value=186, unit="personas", source_name="Nómina",
            status="Aprobado", created_by="cliente@calculatuhuella.local",
        ))
        session.add(ActivityIndicator(
            inventory_id=inv.id, period_start=date(2025, 1, 1), period_end=date(2025, 12, 31),
            indicator_type="Ingresos", value=18_500_000_000, unit="COP", source_name="Estados financieros",
            status="En revisión", created_by="cliente@calculatuhuella.local",
        ))

        session.add_all([
            ReductionAction(
                inventory_id=inv.id, source_id=sources[1].id, title="Optimización de combustión de la caldera",
                description="Ajustar relación aire-combustible, mantenimiento de quemadores y control de eficiencia térmica.",
                baseline_emissions=0, expected_reduction=8.5, investment_cost=28_000_000, annual_savings=42_000_000,
                priority="Alta", responsible="Mantenimiento", target_date=date(2026, 6, 30),
                status="En evaluación", progress_percent=55, created_by="consultor@calculatuhuella.local",
            ),
            ReductionAction(
                inventory_id=inv.id, source_id=sources[0].id, title="Autogeneración solar fotovoltaica",
                description="Instalar un sistema solar para cubrir parte del consumo de la planta Medellín.",
                baseline_emissions=0, expected_reduction=31.0, investment_cost=310_000_000, annual_savings=76_000_000,
                priority="Alta", responsible="Gerencia de operaciones", target_date=date(2027, 3, 31),
                status="Diseño", progress_percent=35, created_by="consultor@calculatuhuella.local",
            ),
            ReductionAction(
                inventory_id=inv.id, source_id=sources[3].id, title="Programa de control de fugas de refrigerantes",
                description="Inventario de equipos, pruebas periódicas de hermeticidad y sustitución progresiva de gases.",
                baseline_emissions=0, expected_reduction=21.5, investment_cost=18_000_000, annual_savings=8_000_000,
                priority="Media", responsible="Mantenimiento", target_date=date(2026, 9, 30),
                status="Identificada", progress_percent=15, created_by="consultor@calculatuhuella.local",
            ),
        ])

        session.add(EmissionTarget(
            inventory_id=inv.id, name="Reducir emisiones absolutas al 2030", metric_type="Absoluta",
            baseline_year=2025, target_year=2030, baseline_value=0, target_value=0, current_value=0,
            unit="tCO₂e", status="Activa", notes="Meta demostrativa: reducción del 20 % frente a la línea base 2025.",
            created_by="consultor@calculatuhuella.local",
        ))

        # Inventario histórico resumido para validar comparación interanual.
        prior = Inventory(
            organization_id=org.id, name="Inventario corporativo 2024", start_date=date(2024, 1, 1), end_date=date(2024, 12, 31),
            objective="Inventario histórico para comparación", base_year=2024, methodology="GHG Protocol + ISO 14064-1",
            methodology_version="GHG Protocol Corporate Standard · ISO 14064-1:2018",
            gwp_version="IPCC AR6 · 100 años", consolidation_approach="Control operacional", materiality_threshold=5,
            status="Cerrado", progress=100, current_stage="Cerrado", notes="Inventario histórico demostrativo.",
            version="1.0", locked=True, closed_by="sistema", closed_at=datetime.now(UTC),
        )
        session.add(prior)
        session.flush()
        for facility in facilities:
            session.add(InventoryFacility(inventory_id=prior.id, facility_id=facility.id, included=True, inclusion_percentage=100))
        prior_sources = [
            EmissionSource(inventory_id=prior.id, facility_id=facilities[0].id, name="Electricidad", scope=2, category="Energía adquirida", progress=100, status="Completado", emissions=52.8),
            EmissionSource(inventory_id=prior.id, facility_id=facilities[0].id, name="Diésel", scope=1, category="Combustión fija", progress=100, status="Completado", emissions=44.9),
            EmissionSource(inventory_id=prior.id, facility_id=facilities[1].id, name="Vehículos", scope=1, category="Combustión móvil", progress=100, status="Completado", emissions=23.6),
            EmissionSource(inventory_id=prior.id, facility_id=facilities[0].id, name="Refrigerantes", scope=1, category="Emisiones fugitivas", progress=100, status="Completado", emissions=79.4),
            EmissionSource(inventory_id=prior.id, facility_id=facilities[0].id, name="Residuos", scope=3, category="Residuos operacionales", progress=100, status="Completado", emissions=61.2),
            EmissionSource(inventory_id=prior.id, facility_id=facilities[1].id, name="Transporte contratado", scope=3, category="Transporte y distribución", progress=100, status="Completado", emissions=15.7),
        ]
        session.add_all(prior_sources)
        session.add_all([
            ActivityIndicator(inventory_id=prior.id, period_start=date(2024, 1, 1), period_end=date(2024, 12, 31), indicator_type="Producción", value=11_840, unit="t", source_name="Registro histórico", status="Aprobado", created_by="sistema"),
            ActivityIndicator(inventory_id=prior.id, period_start=date(2024, 1, 1), period_end=date(2024, 12, 31), indicator_type="Empleados", value=178, unit="personas", source_name="Nómina", status="Aprobado", created_by="sistema"),
            ActivityIndicator(inventory_id=prior.id, period_start=date(2024, 1, 1), period_end=date(2024, 12, 31), indicator_type="Ingresos", value=16_900_000_000, unit="COP", source_name="Estados financieros", status="Aprobado", created_by="sistema"),
        ])

        session.add_all([
            DataRequest(inventory_id=inv.id, source_id=sources[1].id, title="Validar consolidado de diésel", source_name="Diésel", requested_to="Mantenimiento", due_date=date(2026, 8, 15), status="En revisión", instructions="Confirmar consumos mensuales y adjuntar reporte de abastecimiento."),
            DataRequest(inventory_id=inv.id, source_id=sources[5].id, title="Completar transporte contratado", source_name="Transporte contratado", requested_to="Logística", due_date=date(2026, 8, 20), status="Pendiente", instructions="Cargar toneladas, rutas y kilómetros recorridos por proveedor."),
            DataRequest(inventory_id=inv.id, source_id=sources[3].id, title="Subir soporte de refrigerantes", source_name="Refrigerantes", requested_to="Mantenimiento", due_date=date(2026, 8, 22), status="Pendiente", instructions="Adjuntar mantenimientos, recargas y recuperación de gas."),
        ])
        session.add_all([
            ReviewObservation(
                inventory_id=inv.id, source_id=sources[5].id, entity_type="Fuente", entity_label="Transporte contratado",
                title="Completar cobertura anual de transporte contratado",
                description="Solo existen tres periodos y los datos fueron estimados. Completar la actividad anual y adjuntar soporte del proveedor.",
                severity="Mayor", status="Abierta", assigned_to="Logística", due_date=date(2026, 8, 20),
                created_by="revisor@calculatuhuella.local",
            ),
            ReviewObservation(
                inventory_id=inv.id, source_id=sources[3].id, entity_type="Fuente", entity_label="Refrigerantes",
                title="Adjuntar soporte de recargas y recuperación",
                description="El dato anual existe, pero no está relacionado con una evidencia documental.",
                severity="Menor", status="En corrección", assigned_to="Mantenimiento", due_date=date(2026, 8, 22),
                created_by="revisor@calculatuhuella.local", response="Se solicitó el certificado al proveedor de mantenimiento.",
                responded_by="cliente@calculatuhuella.local", responded_at=datetime.now(UTC),
            ),
            ReviewObservation(
                inventory_id=inv.id, source_id=sources[0].id, entity_type="Fuente", entity_label="Electricidad",
                title="Verificar correspondencia factura-periodo",
                description="Validación de muestra documental sobre tres meses del periodo.",
                severity="Informativa", status="Cerrada", assigned_to="Contabilidad",
                created_by="revisor@calculatuhuella.local", response="Se verificaron enero, junio y diciembre.",
                responded_by="cliente@calculatuhuella.local", responded_at=datetime.now(UTC),
                resolution="La muestra coincide con los registros cargados.", resolved_by="revisor@calculatuhuella.local",
                resolved_at=datetime.now(UTC), closed_by="revisor@calculatuhuella.local", closed_at=datetime.now(UTC),
            ),
        ])

        suppliers = [
            Supplier(organization_id=org.id, name="Acero Circular S.A.S.", tax_id="900.111.222-3", sector="Metalmecánico", contact_name="Diana Gómez", contact_email="sostenibilidad@acerocircular.demo", annual_spend_cop=2_450_000_000, strategic=True, risk_level="Alto"),
            Supplier(organization_id=org.id, name="Empaques del Valle S.A.S.", tax_id="901.222.333-4", sector="Empaques", contact_name="Felipe Ríos", contact_email="ambiental@empaquesvalle.demo", annual_spend_cop=980_000_000, strategic=True, risk_level="Medio"),
            Supplier(organization_id=org.id, name="Químicos Andinos Ltda.", tax_id="800.333.444-5", sector="Químicos", contact_name="Laura Peña", contact_email="calidad@quimicosandinos.demo", annual_spend_cop=1_720_000_000, strategic=True, risk_level="Alto"),
            Supplier(organization_id=org.id, name="Servicios Integrales Norte", tax_id="901.444.555-6", sector="Servicios", contact_name="Mateo Díaz", contact_email="operaciones@serviciosnorte.demo", annual_spend_cop=310_000_000, strategic=False, risk_level="Bajo"),
        ]
        session.add_all(suppliers)
        session.flush()
        campaign = SupplierCampaign(
            inventory_id=inv.id, name="Compras estratégicas 2025", category="Bienes y servicios adquiridos",
            due_date=date(2026, 9, 30), status="En curso", methodology="GHG Protocol Scope 3 · categoría 1",
            description="Campaña demostrativa para solicitar factores y huellas específicas a proveedores prioritarios.",
            created_by="consultor@calculatuhuella.local",
        )
        session.add(campaign)
        session.flush()
        supplier_requests = [
            SupplierDataRequest(campaign_id=campaign.id, supplier_id=suppliers[0].id, product_service="Acero laminado", quantity=420, unit="t", spend_cop=2_450_000_000, status="Respondida", due_date=date(2026, 9, 15), access_token="demo-acero-2025", token_expires_at=datetime(2026, 12, 31, tzinfo=UTC), sent_at=datetime.now(UTC), responded_at=datetime.now(UTC)),
            SupplierDataRequest(campaign_id=campaign.id, supplier_id=suppliers[1].id, product_service="Empaques de cartón", quantity=185000, unit="unidad", spend_cop=980_000_000, status="Respondida", due_date=date(2026, 9, 20), access_token="demo-empaques-2025", token_expires_at=datetime(2026, 12, 31, tzinfo=UTC), sent_at=datetime.now(UTC), responded_at=datetime.now(UTC)),
            SupplierDataRequest(campaign_id=campaign.id, supplier_id=suppliers[2].id, product_service="Precursores químicos", quantity=96, unit="t", spend_cop=1_720_000_000, status="Enviada", due_date=date(2026, 9, 25), access_token="demo-quimicos-2025", token_expires_at=datetime(2026, 12, 31, tzinfo=UTC), sent_at=datetime.now(UTC)),
            SupplierDataRequest(campaign_id=campaign.id, supplier_id=suppliers[3].id, product_service="Servicios de mantenimiento", quantity=1, unit="servicio anual", spend_cop=310_000_000, status="Pendiente", due_date=date(2026, 9, 30), access_token="demo-servicios-2025", token_expires_at=datetime(2026, 12, 31, tzinfo=UTC)),
        ]
        session.add_all(supplier_requests)
        session.flush()
        session.add_all([
            SupplierResponse(request_id=supplier_requests[0].id, method="Factor por unidad", activity_value=420, activity_unit="t", emission_factor=710, factor_unit="kg CO2e/t", calculated_emissions_tco2e=298.2, methodology="EPD del producto · cradle-to-gate", boundary="Materias primas, energía y producción hasta puerta de planta", verified=True, quality_level="A", review_status="Aprobado", reviewer_comments="Factor específico y declaración verificada.", reviewed_by="revisor@calculatuhuella.local", reviewed_at=datetime.now(UTC)),
            SupplierResponse(request_id=supplier_requests[1].id, method="Huella total suministrada", activity_value=185000, activity_unit="unidad", reported_emissions_tco2e=42.6, calculated_emissions_tco2e=42.6, methodology="Inventario corporativo asignado por volumen", boundary="Producción y transporte primario", verified=False, quality_level="B", review_status="Aprobado", reviewer_comments="Aceptado con recomendación de mejorar la regla de asignación.", reviewed_by="revisor@calculatuhuella.local", reviewed_at=datetime.now(UTC)),
        ])

        session.flush()
        refresh_progress(session, inv)
        add_audit(session, org.id, "sistema", "CREAR", "Organización", org.name, "Carga inicial de datos demostrativos V0.4")
        add_audit(session, org.id, "sistema", "CONFIGURAR", "Motor de cálculo", "Biblioteca V0.4", "Unidades, GWP, factores y asignaciones iniciales")
        add_audit(session, org.id, "sistema", "CONFIGURAR", "Control profesional", "Flujo V0.5", "Observaciones, decisiones, aprobación, cierre e inmutabilidad")
        session.commit()

        # Import local para evitar dependencia circular durante la definición de modelos.
        from .calculations import recalculate_inventory

        inventory = session.scalar(select(Inventory).where(Inventory.id == inv.id))
        recalculate_inventory(session, inventory)
        session.flush()
        total_emissions = sum(source.emissions for source in inventory.sources if source.included)
        for action in inventory.reduction_actions:
            action.baseline_emissions = action.source.emissions if action.source else total_emissions
        action_profiles = {
            "Optimización de combustión de la caldera": (8, 2026, "Alta", "Bajo"),
            "Autogeneración solar fotovoltaica": (20, 2027, "Media", "Medio"),
            "Programa de control de fugas de refrigerantes": (5, 2026, "Alta", "Bajo"),
        }
        for action in inventory.reduction_actions:
            life, year, feasibility, risk = action_profiles.get(action.title, (5, 2026, "Media", "Medio"))
            action.useful_life_years = life
            action.implementation_year = year
            action.feasibility = feasibility
            action.risk_level = risk
        for target in inventory.targets:
            target.baseline_value = total_emissions
            target.target_value = total_emissions * 0.80
            target.current_value = total_emissions
        scenario = ReductionScenario(
            inventory_id=inventory.id,
            name="Escenario priorizado 2030",
            description="Portafolio demostrativo de medidas priorizadas por costo, reducción y viabilidad.",
            start_year=2026,
            target_year=2030,
            discount_rate=10.0,
            status="En evaluación",
            created_by="consultor@calculatuhuella.local",
        )
        session.add(scenario)
        session.flush()
        for action in inventory.reduction_actions:
            session.add(ReductionScenarioAction(
                scenario_id=scenario.id,
                action_id=action.id,
                included=True,
                implementation_year=action.implementation_year,
                adoption_percent=100.0,
            ))
        session.add(VerificationFinding(
            inventory_id=inventory.id,
            source_id=sources[5].id,
            title="Trazabilidad incompleta del transporte contratado",
            description="La muestra disponible no cubre el periodo completo y utiliza estimaciones sin soporte primario del proveedor.",
            finding_type="Solicitud de información",
            severity="Mayor",
            status="Abierto",
            verifier_email="verificador@calculatuhuella.local",
        ))
        add_audit(session, org.id, "sistema", "CONFIGURAR", "Gestión climática", "Módulos V0.20", "Escenarios comparados, divulgación climática y paquete para comité directivo")
        _ensure_v020_defaults(session)
        _ensure_v021_defaults(session)
        _ensure_v022_defaults(session)
        _ensure_v023_defaults(session)
        _ensure_v024_defaults(session)
        _ensure_v025_defaults(session)
        _ensure_v026_defaults(session)
        _ensure_v027_defaults(session)
        _ensure_v028_defaults(session)
        _ensure_v030_defaults(session)
        _ensure_v031_defaults(session)
        _ensure_v032_defaults(session)
        _ensure_v033_defaults(session)
        _ensure_v034_defaults(session)
        _ensure_v035_defaults(session)
        _ensure_v036_defaults(session)
        _ensure_v037_defaults(session)
        _ensure_v043_defaults(session)
        _ensure_v044_defaults(session)
        _ensure_v045_defaults(session)
        _ensure_v050_defaults(session)
        _ensure_v100_final_defaults(session)
        session.commit()
